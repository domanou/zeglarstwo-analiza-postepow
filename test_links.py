import openpyxl
import pandas as pd

# Ścieżka do pliku
file_path = "data/Dane do analizy.xlsx"

# 1. Odczyt za pomocą openpyxl, aby wyciągnąć hiperłącza
wb = openpyxl.load_workbook(file_path)

print("--- ZNALEZIONE ARKUSZE (ZAWODNICY) ---")
for sheet_name in wb.sheetnames:
    print(f"- {sheet_name}")

    ws = wb[sheet_name]
    print(f"\nPróbka danych i linków z arkusza [{sheet_name}]:")

    # Przeglądamy pierwsze 5 wierszy
    for row in list(ws.iter_rows(values_only=False))[1:6]:
        nazwa_regat = row[4].value  # Kolumna z nazwą regat
        link = (
            row[4].hyperlink.target
            if row[4].hyperlink
            else "Brak linku (zwykły tekst)"
        )
        print(f" Nazwa: {nazwa_regat}")
        print(f" URL:   {link}\n")
